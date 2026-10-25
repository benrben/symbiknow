"""Portable originals exercise real parsers without models, execution or live data."""

from io import BytesIO
from zipfile import ZipFile

import pytest

from symbiknow.ingestion.service import extract_units, index_document, supported
from symbiknow.library import storage
from symbiknow.search.index import document_evidence
from tests.search.test_search import FixtureEmbedder


def office_bytes(kind):
    stream = BytesIO()
    if kind == "docx":
        from docx import Document

        doc = Document()
        doc.add_heading("Atlas ownership", 1)
        doc.add_paragraph("Café originals retain permanent versions.")
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Owner"
        table.cell(0, 1).text = "Mira"
        doc.save(stream)
    else:
        from openpyxl import Workbook

        workbook = Workbook()
        workbook.active.append(["Owner", "Budget"])
        workbook.active.append(["Mira", 42])
        workbook.active.append(["Formula has no cached result", "=1+1"])
        workbook.save(stream)
    return stream.getvalue()


@pytest.mark.parametrize("suffix", ["py", "sh", "svg", "docx", "xlsx", "csv"])
def test_supported_formats_publish_real_text_and_keep_original(tmp_path, suffix):
    storage.init_storage(tmp_path)
    marker = tmp_path / "must-not-run"
    originals = {
        "py": f"# Atlas ownership\nfrom pathlib import Path\nPath({str(marker)!r}).touch()\n# Mira".encode(),
        "sh": f"#!/bin/sh\n# Atlas ownership Mira\ntouch '{marker}'\n".encode(),
        "svg": b'<svg xmlns="http://www.w3.org/2000/svg"><title>Atlas ownership</title>'
               b'<text>Mira <tspan>keeps originals</tspan></text><script>secret script</script>'
               b'<foreignObject><text>untrusted hidden</text></foreignObject>'
               b'<image href="http://127.0.0.1:1/no-request"/></svg>',
        "csv": b"Owner,Budget\nMira,42\n",
    }
    content = office_bytes(suffix) if suffix in {"docx", "xlsx"} else originals[suffix]
    metadata = storage.save_file(tmp_path, content, f"original.{suffix}")
    assert supported(metadata)
    result = index_document(tmp_path, metadata, content, FixtureEmbedder())
    assert result["status"] == "indexed", result
    assert result["searchable_text"] and not result["searchable_visual"]
    evidence = document_evidence(tmp_path, metadata["document_id"], "v1")
    text = "\n".join(row["text"] for row in evidence)
    assert "Mira" in text
    if suffix == "svg":
        assert "keeps originals" in text
        assert "secret script" not in text and "untrusted hidden" not in text
    if suffix == "docx":
        assert "Café originals" in text and "Owner" in text
    if suffix == "xlsx":
        assert "42" in text and "Owner" in text
        assert "=1+1" not in text  # Cached values only; formulas are never evaluated.
    assert not marker.exists()
    assert storage.read_file(tmp_path, metadata["document_id"])[0] == content
    assert len(storage.list_versions(tmp_path, metadata["document_id"])) == 1


@pytest.mark.parametrize("suffix,content", [
    ("svg", b'<!DOCTYPE svg [<!ENTITY x SYSTEM "file:///etc/passwd">]><svg><text>&x;</text></svg>'),
    ("svg", '<!DOCTYPE svg [<!ENTITY x "expansion">]><svg/>'.encode("utf-16")),
    ("docx", b"not an Office archive"),
    ("py", b"\xffnot UTF-8"),
])
def test_invalid_sources_fail_without_losing_original_or_publishing_evidence(tmp_path, suffix, content):
    storage.init_storage(tmp_path)
    metadata = storage.save_file(tmp_path, content, f"invalid.{suffix}")
    result = index_document(tmp_path, metadata, content, FixtureEmbedder())
    assert result["status"] == "failed"
    assert not result["searchable_text"] and not result["searchable_visual"]
    assert document_evidence(tmp_path, metadata["document_id"], "v1") == []
    assert storage.read_file(tmp_path, metadata["document_id"])[0] == content


def test_office_expansion_and_xml_declarations_are_rejected_before_conversion(monkeypatch):
    from symbiknow.ingestion import source_text

    def archive(name, content):
        stream = BytesIO()
        with ZipFile(stream, "w") as output:
            output.writestr(name, content)
        return stream.getvalue()

    with pytest.raises(ValueError, match="DTD"):
        source_text.validate_office(archive("word/document.xml", '<!DOCTYPE x><x/>'.encode("utf-16")))
    with pytest.raises(ValueError, match="structure"):
        source_text.svg_text(b"<svg>" + b"<g>" * 129 + b"</g>" * 129 + b"</svg>")
    monkeypatch.setattr(source_text, "OFFICE_LIMIT", 32)
    with pytest.raises(ValueError, match="budget"):
        source_text.validate_office(archive("xl/worksheets/sheet1.xml", b"x" * 33))
    monkeypatch.setattr(source_text, "TEXT_LIMIT", 4)
    with pytest.raises(ValueError, match="byte limit"):
        source_text.svg_text(b"<svg/>")


@pytest.mark.parametrize("suffix", ["docx", "xlsx"])
def test_real_office_conversion_uses_bounded_isolated_worker(suffix):
    from symbiknow.ingestion.conversion.manager import get_conversion_manager

    manager = get_conversion_manager()
    try:
        units, visual, warnings = extract_units(
            {"filename": f"owned.{suffix}", "mime_type": "application/octet-stream"},
            office_bytes(suffix), FixtureEmbedder(), isolated=True,
        )
        assert "Mira" in "\n".join(row["text"] for row in units)
        assert not visual and not warnings
    finally:
        manager.close()


def test_reconcile_indexes_newly_supported_original_without_retrying_other_unsupported(tmp_path, monkeypatch):
    from symbiknow.api import mcp as server
    from symbiknow.ingestion import service
    from symbiknow.search import index

    monkeypatch.setattr(server, "DATA_DIR", tmp_path)
    monkeypatch.setattr(server, "STORAGE_ONLY", False)
    monkeypatch.delenv("TYPESAFE_API_KEY", raising=False)
    server.application().stop.clear()
    enabled = storage.save_file(tmp_path, b"# Mira owns Atlas source", "source.py")
    unsupported = storage.save_file(tmp_path, b"\x00\xff", "opaque.bin")
    for metadata in (enabled, unsupported):
        index.set_document_status(tmp_path, metadata, status="unsupported")
    calls = []

    def fixture_index(root, metadata, content, **kwargs):
        calls.append(metadata["document_id"])
        return index_document(root, metadata, content, FixtureEmbedder(), **kwargs)

    monkeypatch.setattr(service, "index_document", fixture_index)
    assert server.reconcile_once([enabled["document_id"], unsupported["document_id"]]) == 1
    assert server.reconcile_once([enabled["document_id"], unsupported["document_id"]]) == 0
    assert calls == [enabled["document_id"]]
    assert index.get_document_status(tmp_path, enabled["document_id"])["status"] == "indexed"
    assert index.get_document_status(tmp_path, unsupported["document_id"])["status"] == "unsupported"
    assert len(storage.list_versions(tmp_path, enabled["document_id"])) == 1
